#!/usr/bin/env python3
"""Extrae el contenido completo de un .docx con párrafos NUMERADOS, para citar ubicaciones exactas en el informe.

Uso: python extraer_docx.py ARCHIVO.docx [SALIDA.txt]
Formato: [P012] (estilo) texto   |   [T01 fila 2] celda | celda | celda   |   [PIE] / [ENCABEZADO]
Incluye tablas, pie, encabezado, comentarios y texto oculto/vinculos, que es donde suelen esconderse restos de generación.
Imprime al final un resumen de metadatos del archivo.
"""
import re
import sys
import zipfile

import docx


def main():
    if len(sys.argv) < 2:
        print(__doc__); return 2
    d = docx.Document(sys.argv[1])
    out = []
    for i, p in enumerate(d.paragraphs, 1):
        out.append(f"[P{i:03d}] ({p.style.name}) {p.text}")
    for ti, t in enumerate(d.tables, 1):
        for ri, row in enumerate(t.rows, 1):
            out.append(f"[T{ti:02d} fila {ri}] " + " | ".join(c.text.strip().replace("\n", " / ") for c in row.cells))
    for s in d.sections:
        for etiqueta, parte in (("PIE", s.footer), ("ENCABEZADO", s.header)):
            txt = " ".join(p.text for p in parte.paragraphs).strip()
            if txt:
                out.append(f"[{etiqueta}] {txt}")
    z = zipfile.ZipFile(sys.argv[1])
    if "word/comments.xml" in z.namelist():
        out.append("[COMENTARIOS] " + re.sub(r"<[^>]+>", " ", z.read("word/comments.xml").decode("utf8"))[:2000])
    if re.search(r"<w:vanish", z.read("word/document.xml").decode("utf8")):
        out.append("[AVISO] el documento contiene texto oculto (w:vanish)")
    cp = d.core_properties
    out.append(f"[METADATOS] autor={cp.author!r} ultimo_modificado_por={cp.last_modified_by!r} titulo={cp.title!r} creado={cp.created} comentarios={cp.comments!r}")
    texto = "\n".join(out)
    if len(sys.argv) > 2:
        open(sys.argv[2], "w", encoding="utf-8").write(texto)
        print(f"Escrito {sys.argv[2]} ({len(d.paragraphs)} párrafos, {len(d.tables)} tablas)")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(texto)
    return 0


if __name__ == "__main__":
    sys.exit(main())
